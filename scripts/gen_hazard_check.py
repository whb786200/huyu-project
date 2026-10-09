# -*- coding: utf-8 -*-
"""
隐患排查整改Excel表格生成脚本
从集团安全检查整改回复docx文件自动提取隐患信息，生成带图片的排查表

用法：
  python gen_hazard_check.py <整改回复docx路径> [输出xlsx路径]

示例：
  python gen_hazard_check.py "2026.7.17秦邑集团安全检查问题整改回复 (1).docx"
  python gen_hazard_check.py "整改回复.docx" "输出/隐患排查表.xlsx"
"""

import sys
import os
from datetime import datetime
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter


def extract_images_from_docx(docx_path):
    """从docx中提取所有内嵌图片"""
    doc = Document(docx_path)
    images = []
    img_dir = docx_path + '_extracted_images'
    os.makedirs(img_dir, exist_ok=True)

    for rel in doc.part.rels.values():
        if 'image' in rel.target_ref:
            img_name = os.path.basename(rel.target_ref)
            target_part = rel.target_part
            img_data = target_part.blob
            img_count = len(images) + 1
            img_filename = 'img_%d_%s' % (img_count, img_name)
            img_path = os.path.join(img_dir, img_filename)
            with open(img_path, 'wb') as f:
                f.write(img_data)
            images.append({
                'filename': img_filename,
                'path': img_path,
                'size': len(img_data),
                'ext': os.path.splitext(img_name)[1].lower()
            })

    return images


def extract_issues_from_docx(docx_path):
    """从docx中提取隐患问题描述"""
    doc = Document(docx_path)
    issues = []

    # 遍历段落查找隐患描述
    for p in doc.paragraphs:
        text = p.text.strip()
        if '具体整改如下' in text or '存在问题' in text or '检查发现' in text:
            # 尝试解析编号列表
            lines = text.replace(';', '\n').replace('；', '\n').split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                # 匹配编号格式：1、xxx 或 1.xxx 或 (1)xxx
                for prefix in ['1.', '2.', '3.', '4.', '5.',
                               '1、', '2、', '3、', '4、', '5、',
                               '(1)', '(2)', '(3)', '(4)', '(5)']:
                    if line.startswith(prefix):
                        desc = line[len(prefix):].strip()
                        if desc and len(desc) > 2:
                            issues.append(desc)
                            break

    # 如果段落提取失败，尝试从表格提取
    if not issues:
        for table in doc.tables:
            for row in table.rows:
                cells_text = [cell.text.strip() for cell in row.cells]
                full_text = '|'.join(cells_text)
                if '未' in full_text or '滞后' in full_text or '缺少' in full_text:
                    for cell_text in cells_text:
                        if len(cell_text) > 4 and ('楼' in cell_text or '未' in cell_text or '滞后' in cell_text):
                            if cell_text not in issues:
                                issues.append(cell_text)

    return issues


def create_hazard_check_excel(issues, images, output_path, default_time=None, default_unit=None):
    """创建隐患排查整改Excel表格"""

    wb = Workbook()
    ws = wb.active
    ws.title = '隐患排查整改表'

    # 样式定义
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    cell_font = Font(name='微软雅黑', size=10)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 表头
    headers = ['现场问题照片', '问题描述', '整改措施', '整改完成时间', '整改单位/部门/责任人', '整改后照片']
    col_widths = [18, 22, 26, 16, 22, 18]

    for ci, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(ci)].width = col_widths[ci - 1]

    ws.row_dimensions[1].height = 30

    # 填充数据
    row_height = 120
    default_time = default_time or datetime.now().strftime('%Y年%m月%d日')
    default_unit = default_unit or '陕建集团项目部'

    # 图片配对：前半部分为隐患照片，后半部分为整改后照片
    n_issues = len(issues)
    problem_imgs = images[:n_issues] if len(images) >= n_issues else images
    fixed_imgs = images[n_issues:n_issues * 2] if len(images) >= n_issues * 2 else []

    # 默认责任人建议（根据隐患类型）
    def suggest_responsible(desc):
        desc_lower = desc.lower()
        if '电' in desc or '箱' in desc:
            return default_unit + '/电工'
        elif '消防' in desc or '灭火' in desc:
            return default_unit + '/消防责任人'
        elif '安全' in desc or '带' in desc or '防护' in desc:
            return default_unit + '/安全员'
        else:
            return default_unit + '/安全员'

    def suggest_measure(desc):
        desc_lower = desc.lower()
        if '防护' in desc or '楼梯' in desc:
            return '已按规范要求安装临边防护栏杆'
        elif '巡检' in desc or '记录' in desc:
            return '已更新完善相关巡检记录至最新日期'
        elif '安全带' in desc or '高处' in desc:
            return '已对作业人员进行安全教育，严格要求高处作业必须佩戴安全带'
        elif '灭火' in desc:
            return '已更新灭火器检查记录卡，按要求填写巡检记录'
        else:
            return '已按要求完成整改'

    for ri, issue in enumerate(issues, 2):
        # 文字内容
        ws.cell(row=ri, column=2, value=issue).font = cell_font
        ws.cell(row=ri, column=2).alignment = center_align
        ws.cell(row=ri, column=2).border = thin_border

        measure = suggest_measure(issue)
        ws.cell(row=ri, column=3, value=measure).font = cell_font
        ws.cell(row=ri, column=3).alignment = center_align
        ws.cell(row=ri, column=3).border = thin_border

        ws.cell(row=ri, column=4, value=default_time).font = cell_font
        ws.cell(row=ri, column=4).alignment = center_align
        ws.cell(row=ri, column=4).border = thin_border

        responsible = suggest_responsible(issue)
        ws.cell(row=ri, column=5, value=responsible).font = cell_font
        ws.cell(row=ri, column=5).alignment = center_align
        ws.cell(row=ri, column=5).border = thin_border

        # 图片列边框
        for col in [1, 6]:
            ws.cell(row=ri, column=col).border = thin_border

        ws.row_dimensions[ri].height = row_height

        # 插入问题图片
        if ri - 2 < len(problem_imgs):
            try:
                img = XLImage(problem_imgs[ri - 2]['path'])
                img.width = 150
                ratio = img.height / float(img.width) if img.width > 0 else 0.75
                img.height = int(150 * ratio)
                if img.height > 100:
                    img.height = 100
                    img.width = int(100 / ratio) if ratio > 0 else 150
                ws.add_image(img, '%s%d' % (get_column_letter(1), ri))
            except Exception as e:
                print('[警告] 插入问题图片失败(行%d): %s' % (ri, str(e)))

        # 插入整改后图片
        if ri - 2 < len(fixed_imgs):
            try:
                img2 = XLImage(fixed_imgs[ri - 2]['path'])
                img2.width = 150
                ratio = img2.height / float(img2.width) if img2.width > 0 else 0.75
                img2.height = int(150 * ratio)
                if img2.height > 100:
                    img2.height = 100
                    img2.width = int(100 / ratio) if ratio > 0 else 150
                ws.add_image(img2, '%s%d' % (get_column_letter(6), ri))
            except Exception as e:
                print('[警告] 插入整改图片失败(行%d): %s' % (ri, str(e)))

    wb.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 2:
        print('用法: python gen_hazard_check.py <整改回复docx路径> [输出xlsx路径]')
        print('示例: python gen_hazard_check.py "整改回复.docx"')
        sys.exit(1)

    docx_path = sys.argv[1]
    if not os.path.exists(docx_path):
        print('[错误] 文件不存在: %s' % docx_path)
        sys.exit(1)

    output_path = sys.argv[2] if len(sys.argv) > 2 else '隐患排查整改表.xlsx'

    print('=' * 50)
    print('隐患排查整改Excel表格生成工具')
    print('=' * 50)
    print('[1/3] 正在读取docx文件: %s' % os.path.basename(docx_path))

    issues = extract_issues_from_docx(docx_path)
    print('  提取到 %d 条隐患描述:' % len(issues))
    for i, issue in enumerate(issues, 1):
        print('    %d. %s' % (i, issue))

    print('[2/3] 正在提取内嵌图片...')
    images = extract_images_from_docx(docx_path)
    print('  提取到 %d 张图片' % len(images))
    for img in images:
        print('    - %s (%d KB)' % (img['filename'], img['size'] // 1024))

    print('[3/3] 正在生成Excel表格...')
    result = create_hazard_check_excel(issues, images, output_path)
    print('\n[完成] Excel文件已保存: %s' % result)
    print('共写入 %d 条隐患记录' % len(issues))


if __name__ == '__main__':
    main()
