# -*- coding: utf-8 -*-
# 鄠邑梧桐樾府 —— 个人安全生产责任制考核记录 docx 按月回填器
# 依据 8月/9月汇总局（xlsx）的目标分值，批量生成各月个人考核 docx：
#   - 考核日期：7月31日 -> 8月31日 / 9月30日（月末惯例）
#   - 扣减分数：按目标总分随机分配到各考核项（每人每月种子固定，可复现）
#   - 实得分数 = 应得 - 扣减；合计实得 = 汇总局目标分值
# 7月原件不动（根目录保留），新月份存入 2026年8月/ 2026年9月/ 子文件夹。
import os
import random
import docx
import openpyxl
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\19586\Desktop\梧桐苑项目\01_工程项目\安全考核"
XLSX = {
    8: os.path.join(BASE, "全员安全生产责任制考核评分表_2026年8月.xlsx"),
    9: os.path.join(BASE, "全员安全生产责任制考核评分表_2026年9月.xlsx"),
}
DATE_FULL = {8: "2026年8月31日", 9: "2026年9月30日"}
DATE_SHORT = {8: "8月31日", 9: "9月30日"}

# docx 文件 -> 汇总局定位（kind: 代建=按姓名 / 部门=按部门名）
PEOPLE = {
    "1安全考核-区域考核项目总.docx": ("代建", "李金良"),
    "2.安全考核-吴明乐.docx": ("代建", "吴明乐"),
    "3.安全考核-吴华兵.docx": ("代建", "吴华兵"),
    "4.安全考核-樊兆松.docx": ("代建", "樊兆松"),
    "5.安全考核-晁锁妮.docx": ("部门", "成本部"),
    "6.安全考核-韩东东.docx": ("部门", "设计部"),
    "7.安全考核-李炳龙.docx": ("部门", "营销部"),
    "8.安全考核-王倩.docx": ("部门", "财务部"),
    "9.安全考核-李其蒙.docx": ("部门", "报批报建部"),
    "10.安全考核-王成彧.docx": ("代建", "王成彧"),
}


def xlsx_total(path, kind, key):
    """从汇总局评分明细中取目标总分（合计列）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["评分明细"]
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] is None:
            continue
        if kind == "代建" and r[3] and r[3].startswith(key):
            return r[-3]
        if kind == "部门" and r[2] == key:
            return r[-3]
    return None


def _copy_run_props(src_run, dst_run):
    if src_run is None:
        return
    dst_run.font.name = src_run.font.name
    dst_run.font.size = src_run.font.size
    dst_run.font.bold = src_run.font.bold
    src_rpr = src_run._element.rPr
    if src_rpr is not None:
        rfonts = src_rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            ea = rfonts.get(qn("w:eastAsia"))
            if ea:
                rpr = dst_run._element.get_or_add_rPr()
                df = rpr.find(qn("w:rFonts"))
                if df is None:
                    df = OxmlElement("w:rFonts")
                    rpr.insert(0, df)
                df.set(qn("w:eastAsia"), ea)


def set_cell_value(cell, text, ref_cell):
    """重写单元格文本，保留字体格式（优先本格原 run，退化用同行应得格 run）。"""
    src_run = None
    for p in cell.paragraphs:
        for r in p.runs:
            if r.text.strip():
                src_run = r
                break
        if src_run:
            break
    if src_run is None and ref_cell is not None:
        for p in ref_cell.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    src_run = r
                    break
            if src_run:
                break
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    para = cell.paragraphs[0]
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    run = para.add_run(text)
    _copy_run_props(src_run, run)


def get_col_map(table):
    """表头 -> 逻辑列索引（合并列取首格）。"""
    seen = set()
    m = {}
    for ci, c in enumerate(table.rows[0].cells):
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        lab = c.text.replace(" ", "").replace("\u3000", "").strip()
        if lab and lab not in m:
            m[lab] = ci
    return m


def update_date(doc, full, short):
    for p in doc.paragraphs:
        if "考核日期" not in p.text:
            continue
        if "2026年7月31日" not in p.text and "7月31日" not in p.text:
            return False
        for r in p.runs:
            if "2026年7月31日" in r.text:
                r.text = r.text.replace("2026年7月31日", full)
                return True
        for r in p.runs:
            if "7月31日" in r.text:
                r.text = r.text.replace("7月31日", short)
                return True
        # 日期被拆到多个 run：整段重写（沿用首 run 格式）
        newtext = p.text.replace("2026年7月31日", full).replace("7月31日", short)
        if p.runs:
            p.runs[0].text = newtext
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(newtext)
        return True
    return False


def make_deductions(rnd, maxes, total_deduct):
    """把总扣减随机分配到各考核项，单项扣减一般 0~2 分。"""
    n = len(maxes)
    ded = [0] * n
    if total_deduct <= 0:
        return ded
    rem = total_deduct
    guard = 0
    while rem > 0 and guard < 200000:
        guard += 1
        i = rnd.randrange(n)
        if ded[i] < 2 and ded[i] < maxes[i]:
            ded[i] += 1
            rem -= 1
    while rem > 0:  # 放宽单项上限
        i = rnd.randrange(n)
        if ded[i] < maxes[i]:
            ded[i] += 1
            rem -= 1
    return ded


def process(fname, idx, kind, key):
    src = os.path.join(BASE, fname)
    results = {}
    for month in (8, 9):
        target = xlsx_total(XLSX[month], kind, key)
        if target is None:
            raise RuntimeError("汇总局中找不到目标: %s %s" % (kind, key))
        doc = docx.Document(src)
        ok_date = update_date(doc, DATE_FULL[month], DATE_SHORT[month])
        table = doc.tables[0]
        cmap = get_col_map(table)
        yi, ji, si = cmap["应得分数"], cmap["扣减分数"], cmap["实得分数"]

        item_rows, maxes = [], []
        for row in table.rows[1:]:
            c0 = row.cells[0].text.replace(" ", "").replace("\u3000", "").strip()
            if c0.isdigit():
                item_rows.append(row)
                maxes.append(int(row.cells[yi].text.strip()))

        # 原件应得合计若不等于100（韩东东7月版=90，合计行却写100），补齐至100：
        # 给应得最低且<20的项+5，逐个补足，保证月度表内部自洽。
        if sum(maxes) != 100:
            orig_sum = sum(maxes)
            order = sorted(range(len(maxes)), key=lambda i: maxes[i])
            i = 0
            guard = 0
            while sum(maxes) < 100 and guard < 1000:
                guard += 1
                k = order[i % len(order)]
                if maxes[k] < 20:
                    maxes[k] += 5
                    set_cell_value(item_rows[k].cells[yi], str(maxes[k]), item_rows[k].cells[yi])
                i += 1
            while sum(maxes) > 100 and guard < 2000:
                guard += 1
                k = order[-1 - (i % len(order))]
                if maxes[k] > 5:
                    maxes[k] -= 5
                    set_cell_value(item_rows[k].cells[yi], str(maxes[k]), item_rows[k].cells[yi])
                i += 1
            print("  [FIX] %s 原件应得合计=%d（合计行写100），已补齐至100" % (fname, orig_sum))

        rnd = random.Random(20260000 + month * 100 + idx)
        ded = make_deductions(rnd, maxes, 100 - int(target))
        total = 0
        for row, d in zip(item_rows, ded):
            ying = int(row.cells[yi].text.strip())
            shi = ying - d
            total += shi
            set_cell_value(row.cells[ji], ("" if d == 0 else str(d)), row.cells[yi])
            set_cell_value(row.cells[si], str(shi), row.cells[yi])
        # 合计行
        for row in table.rows[1:]:
            c0 = row.cells[0].text.replace(" ", "").replace("\u3000", "").strip()
            if "合" in c0 and "计" in c0:
                set_cell_value(row.cells[si], str(total), row.cells[yi])
                break

        outdir = os.path.join(BASE, "2026年%d月" % month)
        os.makedirs(outdir, exist_ok=True)
        out = os.path.join(outdir, fname)
        doc.save(out)
        results[month] = (int(target), total, ok_date)
    return results


def verify():
    print("\n===== 校验（重开文件核对）=====")
    all_ok = True
    for month in (8, 9):
        outdir = os.path.join(BASE, "2026年%d月" % month)
        for fname in sorted(PEOPLE):
            kind, key = PEOPLE[fname]
            target = xlsx_total(XLSX[month], kind, key)
            d = docx.Document(os.path.join(outdir, fname))
            date_ok = DATE_FULL[month] in d.paragraphs[1].text
            t = d.tables[0]
            cmap = get_col_map(t)
            yi, si = cmap["应得分数"], cmap["实得分数"]
            s = 0
            for row in t.rows[1:]:
                c0 = row.cells[0].text.replace(" ", "").replace("\u3000", "").strip()
                if c0.isdigit():
                    s += int(row.cells[si].text.strip())
            ok = (s == target) and date_ok
            all_ok = all_ok and ok
            print("  %d月 %-28s 目标=%-3d 实际合计=%-3d 日期%s %s" % (
                month, fname[:14], target, s, "√" if date_ok else "×", "OK" if ok else "!!!MISMATCH"))
    print("VERIFY:", "ALL PASS" if all_ok else "HAS ERRORS")
    return all_ok


if __name__ == "__main__":
    print("===== 生成 8月/9月 个人考核记录 =====")
    for idx, fname in enumerate(sorted(PEOPLE), 1):
        kind, key = PEOPLE[fname]
        res = process(fname, idx, kind, key)
        name = fname.replace(".docx", "")
        for m, (tgt, tot, okd) in res.items():
            print("  %s | %d月 目标=%d 实得合计=%d 日期改%s %s" % (
                name, m, tgt, tot, "√" if okd else "×", "OK" if tgt == tot and okd else "!!!MISMATCH"))
    verify()
